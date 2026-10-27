"""Read-only implementation and packaging checks; not a new official-source audit."""
import hashlib,json,math,pathlib,zipfile
import openpyxl
from pypdf import PdfReader
ROOT=pathlib.Path(__file__).resolve().parents[1]
book=ROOT/'workbooks/KOHC_Financial_Analysis_FINAL_SUBMISSION.xlsx';audit=ROOT/'workbooks/KOHC_Financial_Analysis_Audit.xlsx'
before={str(p):hashlib.sha256(p.read_bytes()).hexdigest()for p in [book,audit]}
w=openpyxl.load_workbook(book,data_only=True);f=openpyxl.load_workbook(book)
cos=['KOHC','LUCK','DGKC','MLCF','FCCL'];rows=[111,112,113,114,116,117,118,119,120,121,123,124,125,126,128,129,130,131,132,133,134,136,137,138,139,140,141]
same=lambda a,b:math.isclose(a,b,rel_tol=1e-10,abs_tol=1e-10)if isinstance(a,(int,float))and isinstance(b,(int,float))else a==b
checks={}
for co in cos:
 s=w[co]
 assert s.column_dimensions['I'].hidden and s['I7'].value=='FY2020*'
 for year in range(2021,2026):
  c=year-2017;pc=9 if year==2021 else c-1;v=lambda r:s.cell(r,c).value;p=lambda r:s.cell(r,pc).value
  cash=154 if co=='MLCF'else 157;ev=157 if co=='MLCF'else 160;inv=(v(54)+p(54))/2;receivables=(v(20)+p(20))/2
  pairs=[(v(24),v(48)),(v(24)-v(54),v(48)),(v(cash),v(48)),(v(24)-v(48),1),(-v(65),inv),(v(64),receivables),(365*inv,-v(65)),(365*receivables,v(64)),(v(64),v(9)),(v(64),v(25)),(v(49),v(25)),(v(55),v(29)),(v(55),v(55)+v(29)),(v(82),v(85)),(v(66),v(64)),(v(82),v(64)),(v(79),v(64)),(v(79),v(25)),(v(79),v(29)),(v(88),v(58)),(v(82),v(25)),(v(79),v(104)),(v(29),v(103)),(v(105),v(79)/v(104)),(v(105),v(29)/v(103)),(v(106),v(105)),(v(ev),v(83))]
  for row,(num,den)in zip(rows,pairs):
   expected='n/m'if num<0 or den<=0 else num/den
   assert f[co].cell(row,c).data_type=='f',(co,year,row,'missing formula')
   assert same(expected,v(row)),(co,year,row,expected,v(row));checks[co,year,row]=v(row)
 for r in [110,115,122,127,135]:assert all(f[co].cell(r,c).value is None for c in range(2,9))
 assert not any(isinstance(c.value,str)and('vertical analysis'in c.value.lower()or'horizontal analysis'in c.value.lower())for row in s for c in row)
assert len(checks)==675
meanrows=[r[0].row for r in w['Industry Averages']if r[0].value=='Five-company average'];assert len(meanrows)==27
for ratio_row,meanrow in zip(rows,meanrows):
 for year in range(2021,2026):
  vals=[checks[co,year,ratio_row]for co in cos if isinstance(checks[co,year,ratio_row],(int,float))]
  assert same(sum(vals)/len(vals),w['Industry Averages'].cell(meanrow,year-2017).value)
  assert f['Industry Averages'].cell(meanrow,year-2017).data_type=='f'
assert len(f['Charts']._charts)==54 and sum(len(c.series)for c in f['Charts']._charts)==189
with zipfile.ZipFile(book)as z:
 charts=[n for n in z.namelist()if n.startswith('xl/charts/chart')and n.endswith('.xml')];assert len(charts)==54
 assert all(b'#REF!'not in z.read(n)for n in z.namelist()if n.endswith('.xml'))
errors=[(s.title,c.coordinate,c.value)for s in w for row in s for c in row if c.data_type=='e']
assert len(errors)==11 and all(s=='Charts'and val=='#N/A'for s,addr,val in errors)
obs=json.loads((ROOT/'evidence/Dividend_Official_Source_Check.json').read_text())['observations'];assert len(obs)==25
for x in obs:
 c=x['year']-2017;s=w[x['co']]
 for row,key in [(106,'dps'),(105,'price'),(140,'yield')]:assert same(s.cell(row,c).value,x[key])
assert sum(x['dps']==0 for x in obs)==17
for p in [book,audit]:assert hashlib.sha256(p.read_bytes()).hexdigest()==before[str(p)]
pdf=PdfReader(ROOT/'docs/KOHC_Viva_Guide_Updated_10_Pages.pdf');assert len(pdf.pages)==10
manifest=json.loads((ROOT/'evidence/FILE_MANIFEST.json').read_text())
for path,entry in manifest.items():
 p=ROOT/path;assert p.stat().st_size==entry['bytes']and hashlib.sha256(p.read_bytes()).hexdigest()==entry['sha256'],path
result={'ratio_results_reproduced_from_existing_inputs':675,'sample_means_verified':135,'charts':54,'series':189,'dividend_observations_checked_against_existing_official_record':25,'numeric_zero_dividends':17,'intentional_chart_gap_errors':11,'unexpected_errors':0,'updated_viva_pages':10,'packaged_file_hashes_verified':len(manifest),'workbooks_unchanged':True,'qualification':'KOHC accrued-markup classification and18 published-value differences remain; arithmetic reproduction is not resolution of these source/classification issues.'}
print(json.dumps(result,indent=2))
